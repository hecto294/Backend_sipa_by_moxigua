import csv
import io
from datetime import date, datetime
from typing import List, Dict, Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet


def _serializar(valor: Any) -> Any:
    if isinstance(valor, (date, datetime)):
        return valor.isoformat()
    return valor


def _build_excel(headers: List[str], rows: List[Dict[str, Any]], title: str) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal="center")

    for col_idx, header in enumerate(headers, start=1):
        c = ws.cell(row=2, column=col_idx, value=header)
        c.font = Font(bold=True)
        c.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")

    for row_idx, row_data in enumerate(rows, start=3):
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=row_idx, column=col_idx, value=_serializar(row_data.get(header, "")))

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _build_pdf(headers: List[str], rows: List[Dict[str, Any]], title: str) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1*cm,
        rightMargin=1*cm,
        topMargin=1.5*cm,
        bottomMargin=1*cm,
    )
    elements = []
    styles = getSampleStyleSheet()
    elements.append(Paragraph(title, styles["Title"]))
    elements.append(Spacer(1, 12))

    table_data = [headers]
    for row in rows:
        table_data.append([_serializar(row.get(h, "")) for h in headers])

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 10),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)
    return buffer


def _build_csv(headers: List[str], rows: List[Dict[str, Any]]) -> io.StringIO:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in rows:
        writer.writerow({h: _serializar(row.get(h, "")) for h in headers})
    output.seek(0)
    return output


def generar_reporte(
    titulo: str,
    headers: List[str],
    rows: List[Dict[str, Any]],
    formato: str,
) -> tuple[Any, str, str]:
    """
    Genera el reporte en el formato solicitado.
    Retorna (buffer, content_type, extension).
    """
    if formato == "xlsx":
        buffer = _build_excel(headers, rows, titulo)
        return buffer, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    elif formato == "pdf":
        buffer = _build_pdf(headers, rows, titulo)
        return buffer, "application/pdf", "pdf"
    elif formato == "csv":
        buffer = _build_csv(headers, rows)
        return buffer, "text/csv", "csv"
    else:
        raise ValueError("Formato no soportado")