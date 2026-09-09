import csv
import io
from datetime import date
from typing import Any, Dict, List, Tuple

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.security import get_password_hash
from app.models import (
    Usuario,
    Ficha,
    AsignacionInstructorFicha,
    ProcesoEtapaProductiva,
    Rol,
    ProgramaFormacion,
    ModalidadEP,
    Empresa,
    CoordinadorEmpresa,
)
from app.modules.importacion.schemas import RowError


# Definición de los tipos de importación soportados
TIPOS_IMPORTACION = {
    "usuarios": "Usuarios (aprendices, instructores, coordinadores)",
    "fichas": "Fichas de formación",
    "asignaciones": "Asignaciones instructor-ficha",
    "procesos": "Procesos de etapa productiva",
}


def leer_archivo(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Lee un archivo CSV o Excel y lo convierte en una lista de diccionarios.
    Retorna la lista de filas.
    """
    if filename.lower().endswith(".csv"):
        content = file_bytes.decode("utf-8")
        reader = csv.DictReader(io.StringIO(content))
        return [row for row in reader]

    elif filename.lower().endswith((".xlsx", ".xls")):
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(h) if h else "" for h in rows[0]]
        result = []
        for row in rows[1:]:
            if all(cell is None or cell == "" for cell in row):
                continue  # ignorar filas vacías
            row_dict = {
                headers[i]: row[i] if i < len(row) else None
                for i in range(len(headers))
            }
            result.append(row_dict)
        return result

    else:
        raise ValueError("Formato de archivo no soportado. Use CSV o Excel (.xlsx)")


# ============================================================
# Validadores de fila por tipo de importación
# ============================================================

def validar_fila_usuario(data: Dict[str, Any]) -> Tuple[bool, str, Dict[str, Any]]:
    """Valida una fila de usuario. Retorna (ok, mensaje_error, datos_procesados)."""
    campos_requeridos = ["nombre", "apellido", "email", "password_hash", "rol_id"]
    for campo in campos_requeridos:
        if not data.get(campo):
            return False, f"Campo requerido faltante: {campo}", None

    email = data["email"].strip()
    if "@" not in email or "." not in email:
        return False, f"Email inválido: {email}", None

    try:
        rol_id = int(data["rol_id"])
    except (ValueError, TypeError):
        return False, f"rol_id no es un número: {data['rol_id']}", None

    # El password se recibe plano y se hashea
    password_plano = data["password_hash"]
    if len(password_plano) < 8:
        return False, "La contraseña debe tener al menos 8 caracteres", None

    datos_procesados = {
        "nombre": data["nombre"].strip(),
        "apellido": data["apellido"].strip(),
        "email": email,
        "password_hash": get_password_hash(password_plano),
        "documento_identidad": data.get("documento_identidad") or None,
        "telefono": data.get("telefono") or None,
        "rol_id": rol_id,
        "is_active": True,
    }
    return True, "", datos_procesados


def validar_fila_ficha(data: Dict[str, Any]) -> Tuple[bool, str, Dict[str, Any]]:
    """Valida una fila de ficha. Retorna (ok, mensaje_error, datos_procesados)."""
    campos_requeridos = ["programa_id", "numero_ficha", "fecha_inicio", "fecha_fin"]
    for campo in campos_requeridos:
        if not data.get(campo):
            return False, f"Campo requerido faltante: {campo}", None

    try:
        programa_id = int(data["programa_id"])
    except (ValueError, TypeError):
        return False, f"programa_id no es un número: {data['programa_id']}", None

    try:
        fecha_inicio = date.fromisoformat(data["fecha_inicio"])
        fecha_fin = date.fromisoformat(data["fecha_fin"])
        if fecha_fin < fecha_inicio:
            return False, "fecha_fin no puede ser anterior a fecha_inicio", None
    except ValueError:
        return False, "Fechas inválidas, use formato YYYY-MM-DD", None

    datos_procesados = {
        "programa_id": programa_id,
        "numero_ficha": data["numero_ficha"].strip(),
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "is_active": True,
    }
    return True, "", datos_procesados


def validar_fila_asignacion(data: Dict[str, Any]) -> Tuple[bool, str, Dict[str, Any]]:
    """Valida una fila de asignación. Retorna (ok, mensaje_error, datos_procesados)."""
    campos_requeridos = ["ficha_id", "instructor_id"]
    for campo in campos_requeridos:
        if not data.get(campo):
            return False, f"Campo requerido faltante: {campo}", None

    try:
        ficha_id = int(data["ficha_id"])
        instructor_id = int(data["instructor_id"])
    except (ValueError, TypeError):
        return False, "ficha_id e instructor_id deben ser números", None

    datos_procesados = {
        "ficha_id": ficha_id,
        "instructor_id": instructor_id,
        "estado_asignacion": "ACTIVA",
        "is_active": True,
    }
    return True, "", datos_procesados


def validar_fila_proceso(data: Dict[str, Any]) -> Tuple[bool, str, Dict[str, Any]]:
    """Valida una fila de proceso. Retorna (ok, mensaje_error, datos_procesados)."""
    campos_requeridos = ["aprendiz_id", "ficha_id", "modalidad_id", "instructor_id",
                         "fecha_inicio", "fecha_fin"]
    for campo in campos_requeridos:
        if not data.get(campo):
            return False, f"Campo requerido faltante: {campo}", None

    try:
        aprendiz_id = int(data["aprendiz_id"])
        ficha_id = int(data["ficha_id"])
        modalidad_id = int(data["modalidad_id"])
        instructor_id = int(data["instructor_id"])
    except (ValueError, TypeError):
        return False, "IDs deben ser numéricos", None

    try:
        fecha_inicio = date.fromisoformat(data["fecha_inicio"])
        fecha_fin = date.fromisoformat(data["fecha_fin"])
        if fecha_fin < fecha_inicio:
            return False, "fecha_fin no puede ser anterior a fecha_inicio", None
    except ValueError:
        return False, "Fechas inválidas, use formato YYYY-MM-DD", None

    # Convertir opcionales empresa_id y coordinador_empresa_id
    empresa_id = None
    coordinador_empresa_id = None

    raw_empresa = data.get("empresa_id")
    if raw_empresa not in (None, "", "None", "null"):
        try:
            empresa_id = int(raw_empresa)
        except (ValueError, TypeError):
            return False, "empresa_id debe ser numérico o vacío", None

    raw_coord = data.get("coordinador_empresa_id")
    if raw_coord not in (None, "", "None", "null"):
        try:
            coordinador_empresa_id = int(raw_coord)
        except (ValueError, TypeError):
            return False, "coordinador_empresa_id debe ser numérico o vacío", None

    if (empresa_id is None) != (coordinador_empresa_id is None):
        return False, "empresa_id y coordinador_empresa_id deben venir juntos o ambos vacíos", None

    datos_procesados = {
        "aprendiz_id": aprendiz_id,
        "ficha_id": ficha_id,
        "modalidad_id": modalidad_id,
        "instructor_id": instructor_id,
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "empresa_id": empresa_id,
        "coordinador_empresa_id": coordinador_empresa_id,
        "estado": "ACTIVO",
        "estado_sofia": "PENDIENTE",
        "is_active": True,
    }
    return True, "", datos_procesados


# Mapeo de tipo_importacion a su función de validación
VALIDACIONES = {
    "usuarios": validar_fila_usuario,
    "fichas": validar_fila_ficha,
    "asignaciones": validar_fila_asignacion,
    "procesos": validar_fila_proceso,
}


def importar_datos(
    db: Session,
    tipo: str,
    filas: List[Dict[str, Any]]
) -> Tuple[int, int, List[RowError]]:
    """
    Procesa las filas y las inserta en la base de datos.
    Usa transacciones anidadas (savepoints) para aislar cada fila.
    Retorna (filas_insertadas, filas_con_error, errores).
    """
    if tipo not in VALIDACIONES:
        raise ValueError(f"Tipo de importación no soportado: {tipo}")

    filas_insertadas = 0
    filas_con_error = 0
    errores: List[RowError] = []
    fila_num = 1

    for fila in filas:
        fila_num += 1
        valida, msg_error, datos = VALIDACIONES[tipo](fila)
        if not valida:
            filas_con_error += 1
            errores.append(RowError(fila=fila_num, mensaje=msg_error, datos=fila))
            continue

        try:
            # Punto de guardado para revertir solo esta fila si falla
            with db.begin_nested():
                nuevo = _crear_instancia_validada(db, tipo, datos)
                db.add(nuevo)

            filas_insertadas += 1

        except Exception as e:
            # El savepoint revirtió automáticamente esta fila
            filas_con_error += 1
            errores.append(RowError(fila=fila_num, mensaje=str(e), datos=fila))

    # Commit final para consolidar todas las filas exitosas
    db.commit()
    return filas_insertadas, filas_con_error, errores


def _crear_instancia_validada(db: Session, tipo: str, datos: Dict[str, Any]):
    """
    Crea la instancia SQLAlchemy correspondiente al tipo de importación,
    validando previamente las claves foráneas y reglas de negocio.
    Lanza ValueError con mensaje claro si algo falla.
    """
    if tipo == "usuarios":
        rol_id = datos.get("rol_id")
        if not db.query(Rol).filter(Rol.id == rol_id, Rol.is_active == True).first():
            raise ValueError(f"Rol ID {rol_id} no existe o está inactivo")

        # Validar duplicidad
        if db.query(Usuario).filter(
            (Usuario.email == datos["email"]) |
            (Usuario.documento_identidad == datos.get("documento_identidad"))
        ).first():
            raise ValueError("Email o documento ya registrado")

        return Usuario(**datos)

    elif tipo == "fichas":
        programa_id = datos.get("programa_id")
        if not db.query(ProgramaFormacion).filter(
            ProgramaFormacion.id == programa_id,
            ProgramaFormacion.is_active == True
        ).first():
            raise ValueError(f"Programa ID {programa_id} no existe o está inactivo")

        if db.query(Ficha).filter(Ficha.numero_ficha == datos["numero_ficha"]).first():
            raise ValueError("Número de ficha duplicado")

        return Ficha(**datos)

    elif tipo == "asignaciones":
        ficha_id = datos.get("ficha_id")
        instructor_id = datos.get("instructor_id")

        if not db.query(Ficha).filter(Ficha.id == ficha_id, Ficha.is_active == True).first():
            raise ValueError(f"Ficha ID {ficha_id} no existe o está inactiva")

        if not db.query(Usuario).filter(
            Usuario.id == instructor_id,
            Usuario.rol_id == 3,
            Usuario.is_active == True
        ).first():
            raise ValueError(f"Instructor ID {instructor_id} no es un instructor activo")

        if db.query(AsignacionInstructorFicha).filter(
            AsignacionInstructorFicha.ficha_id == ficha_id,
            AsignacionInstructorFicha.instructor_id == instructor_id,
            AsignacionInstructorFicha.estado_asignacion == "ACTIVA"
        ).first():
            raise ValueError("Asignación ya existe")

        return AsignacionInstructorFicha(**datos)

    elif tipo == "procesos":
        aprendiz_id = datos.get("aprendiz_id")
        ficha_id = datos.get("ficha_id")
        modalidad_id = datos.get("modalidad_id")
        instructor_id = datos.get("instructor_id")
        empresa_id = datos.get("empresa_id")
        coordinador_id = datos.get("coordinador_empresa_id")

        if not db.query(Usuario).filter(
            Usuario.id == aprendiz_id,
            Usuario.rol_id == 4,
            Usuario.is_active == True
        ).first():
            raise ValueError(f"Aprendiz ID {aprendiz_id} no es un aprendiz activo")

        if not db.query(Ficha).filter(Ficha.id == ficha_id, Ficha.is_active == True).first():
            raise ValueError(f"Ficha ID {ficha_id} no existe o está inactiva")

        if not db.query(ModalidadEP).filter(
            ModalidadEP.id == modalidad_id,
            ModalidadEP.is_active == True
        ).first():
            raise ValueError(f"Modalidad ID {modalidad_id} no existe o está inactiva")

        if not db.query(Usuario).filter(
            Usuario.id == instructor_id,
            Usuario.rol_id == 3,
            Usuario.is_active == True
        ).first():
            raise ValueError(f"Instructor ID {instructor_id} no es un instructor activo")

        # Validar empresa y coordinador si vienen
        if empresa_id is not None:
            if not db.query(Empresa).filter(Empresa.id == empresa_id, Empresa.is_active == True).first():
                raise ValueError(f"Empresa ID {empresa_id} no existe o está inactiva")
            if coordinador_id is None or not db.query(CoordinadorEmpresa).filter(
                CoordinadorEmpresa.id == coordinador_id,
                CoordinadorEmpresa.empresa_id == empresa_id,
                CoordinadorEmpresa.is_active == True
            ).first():
                raise ValueError("Coordinador de empresa no válido para la empresa indicada")
        elif coordinador_id is not None:
            raise ValueError("Se indicó coordinador_empresa_id sin empresa_id")

        return ProcesoEtapaProductiva(**datos)

    else:
        raise ValueError(f"Tipo no soportado: {tipo}")